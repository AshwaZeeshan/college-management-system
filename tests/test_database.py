"""Regression checks using temporary databases, independent of college.db."""

import csv
from pathlib import Path
import sqlite3
import sys
import tempfile
import unittest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import auth
import db
import reports
from openpyxl import load_workbook


class DatabaseTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.old_db_path = db.DB_PATH
        self.old_reports_dir = reports.REPORTS_DIR
        self.addCleanup(self.restore_paths)
        db.DB_PATH = str(Path(self.temp.name) / "test.db")
        reports.REPORTS_DIR = self.temp.name

    def restore_paths(self):
        db.DB_PATH = self.old_db_path
        reports.REPORTS_DIR = self.old_reports_dir

    def test_fresh_database_and_repeat_startup(self):
        db.init_db()
        with db.get_connection() as conn:
            before = {table: conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
                      for table in ("users", "students", "professors", "lab_attendants", "courses")}
            self.assertEqual(conn.execute("PRAGMA integrity_check").fetchone()[0], "ok")
            self.assertEqual(conn.execute("PRAGMA foreign_key_check").fetchall(), [])
        conn.close()
        self.assertGreater(before["students"], 0)
        db.init_db()
        with db.get_connection() as conn:
            after = {table: conn.execute(f"SELECT COUNT(*) FROM {table}").fetchone()[0]
                     for table in before}
        conn.close()
        self.assertEqual(before, after)

    def test_migration_handles_missing_and_existing_tables(self):
        conn = sqlite3.connect(":memory:")
        self.addCleanup(conn.close)
        db._migrate_existing_db(conn.cursor())
        conn.execute("CREATE TABLE users (user_id INTEGER PRIMARY KEY)")
        conn.execute("INSERT INTO users VALUES (7)")
        db._migrate_existing_db(conn.cursor())
        db._migrate_existing_db(conn.cursor())
        self.assertEqual(conn.execute("SELECT user_id, is_active FROM users").fetchone(), (7, 1))

    def test_demo_authentication_and_disabled_account(self):
        db.init_db()
        for email, password, role in (
            ("admin@college.edu", "Admin@1234", "admin"),
            ("student101@college.edu", "Sample@1234", "student"),
            ("prof201@college.edu", "Sample@1234", "professor"),
            ("lab301@college.edu", "Sample@1234", "lab_attendant"),
        ):
            with self.subTest(role=role):
                user, role_id = auth.login(email, password)
                self.assertEqual(user["role"], role)
                if role != "admin":
                    self.assertIsNotNone(role_id)
        with self.assertRaises(ValueError):
            auth.login("admin@college.edu", "WrongPassword1!")
        with db.get_connection() as conn:
            conn.execute("UPDATE users SET is_active=0 WHERE email='student101@college.edu'")
        conn.close()
        with self.assertRaises(auth.AccountDisabledError):
            auth.login("student101@college.edu", "Sample@1234")

    def test_all_report_formats(self):
        db.init_db()
        generators = (
            reports.generate_students_report, reports.generate_professors_report,
            reports.generate_courses_report, reports.generate_enrollments_report,
            reports.generate_lab_report, reports.generate_summary_report,
        )
        for generate in generators:
            for fmt in ("pdf", "excel", "csv"):
                with self.subTest(report=generate.__name__, format=fmt):
                    path = Path(generate(fmt))
                    self.assertGreater(path.stat().st_size, 0)
                    if fmt == "pdf":
                        self.assertTrue(path.read_bytes().startswith(b"%PDF-"))
                    elif fmt == "excel":
                        book = load_workbook(path, read_only=True)
                        try:
                            self.assertGreater(book.active.max_row, 1)
                        finally:
                            book.close()
                    else:
                        with path.open(encoding="utf-8-sig", newline="") as handle:
                            self.assertGreater(len(list(csv.reader(handle))), 1)


if __name__ == "__main__":
    unittest.main()
