"""
reports.py  –  Export engine for College Management System.

Supports three formats:
  • PDF   – via ReportLab (professional layout)
  • Excel – via openpyxl (styled .xlsx)
  • CSV   – via Python stdlib csv module

Every public generate_* function accepts a `fmt` parameter:
    fmt = 'pdf' | 'excel' | 'csv'
and returns the saved file path.
"""

import os
import csv
import datetime
from db import get_connection

# ── ReportLab ───────────────────────────────────────────────────────────────
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import (
    SimpleDocTemplate, Table, TableStyle,
    Paragraph, Spacer, HRFlowable,
)
from reportlab.lib.enums import TA_CENTER, TA_LEFT

# ── openpyxl ────────────────────────────────────────────────────────────────
import openpyxl
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, GradientFill
)
from openpyxl.utils import get_column_letter

# ── Directories ─────────────────────────────────────────────────────────────
REPORTS_DIR = os.path.join(os.path.dirname(__file__), "reports")
os.makedirs(REPORTS_DIR, exist_ok=True)

# ── Colour palette ───────────────────────────────────────────────────────────
PRIMARY   = colors.HexColor("#1a2744")
SECONDARY = colors.HexColor("#2980b9")
ACCENT    = colors.HexColor("#27ae60")
LIGHT_BG  = colors.HexColor("#eaf0fb")
WHITE     = colors.white
GREY_CLR  = colors.HexColor("#7f8c8d")

# Excel hex colours (no #)
XL_NAVY   = "1a2744"
XL_BLUE   = "2980b9"
XL_GREEN  = "27ae60"
XL_LIGHT  = "eaf0fb"
XL_WHITE  = "FFFFFF"
XL_GREY   = "f4f6f7"


# ════════════════════════════════════════════════════════════════════════════
#  INTERNAL HELPERS
# ════════════════════════════════════════════════════════════════════════════

def _save_path(name: str, fmt: str) -> str:
    ext_map = {"pdf": "pdf", "excel": "xlsx", "csv": "csv"}
    ts  = datetime.datetime.now().strftime("%Y%m%d_%H%M%S")
    return os.path.join(REPORTS_DIR, f"{name}_{ts}.{ext_map[fmt]}")


# ── PDF helpers ──────────────────────────────────────────────────────────── #

def _pdf_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("RTitle", fontSize=18, fontName="Helvetica-Bold",
                               textColor=PRIMARY, alignment=TA_CENTER, spaceAfter=4))
    styles.add(ParagraphStyle("RSub",   fontSize=10, fontName="Helvetica",
                               textColor=GREY_CLR, alignment=TA_CENTER, spaceAfter=10))
    return styles


def _pdf_intro(styles, title, subtitle):
    now = datetime.datetime.now().strftime("%d %B %Y  %H:%M")
    return [
        Paragraph("College Management System", styles["RTitle"]),
        Paragraph(title, styles["RSub"]),
        Paragraph(f"Generated: {now}  ·  {subtitle}", styles["RSub"]),
        HRFlowable(width="100%", thickness=1.5, color=SECONDARY),
        Spacer(1, 0.4 * cm),
    ]


def _pdf_table_style():
    return TableStyle([
        ("BACKGROUND",    (0, 0), (-1, 0),  PRIMARY),
        ("TEXTCOLOR",     (0, 0), (-1, 0),  WHITE),
        ("FONTNAME",      (0, 0), (-1, 0),  "Helvetica-Bold"),
        ("FONTSIZE",      (0, 0), (-1, 0),  8),
        ("TOPPADDING",    (0, 0), (-1, 0),  7),
        ("BOTTOMPADDING", (0, 0), (-1, 0),  7),
        ("ROWBACKGROUNDS",(0, 1), (-1, -1), [WHITE, LIGHT_BG]),
        ("FONTNAME",      (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE",      (0, 1), (-1, -1), 7),
        ("ALIGN",         (0, 0), (-1, -1), "CENTER"),
        ("VALIGN",        (0, 0), (-1, -1), "MIDDLE"),
        ("GRID",          (0, 0), (-1, -1), 0.35, colors.HexColor("#c0d0e0")),
        ("TOPPADDING",    (0, 1), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 1), (-1, -1), 4),
    ])


def _build_pdf(path, title, subtitle, headers, rows, col_widths):
    doc  = SimpleDocTemplate(path, pagesize=A4,
                             leftMargin=1.2*cm, rightMargin=1.2*cm,
                             topMargin=1.5*cm,  bottomMargin=1.5*cm)
    sty  = _pdf_styles()
    data = [headers] + [[str(c) if c is not None else "" for c in r] for r in rows]
    tbl  = Table(data, colWidths=col_widths, repeatRows=1)
    tbl.setStyle(_pdf_table_style())
    doc.build(_pdf_intro(sty, title, subtitle) + [tbl])


# ── Excel helpers ────────────────────────────────────────────────────────── #

_thin = Side(style="thin", color="C0D0E0")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _xl_header_style():
    return {
        "font":      Font(name="Calibri", bold=True, color=XL_WHITE, size=10),
        "fill":      PatternFill("solid", fgColor=XL_NAVY),
        "alignment": Alignment(horizontal="center", vertical="center"),
        "border":    _border,
    }


def _xl_data_style(even: bool):
    return {
        "font":      Font(name="Calibri", size=9),
        "fill":      PatternFill("solid", fgColor=XL_LIGHT if even else XL_WHITE),
        "alignment": Alignment(horizontal="center", vertical="center"),
        "border":    _border,
    }


def _xl_apply(cell, style_dict):
    for attr, val in style_dict.items():
        setattr(cell, attr, val)


def _build_excel(path, sheet_name, title, headers, rows):
    wb  = openpyxl.Workbook()
    ws  = wb.active
    ws.title = sheet_name

    # Title row
    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1,   end_column=len(headers))
    tc = ws.cell(row=1, column=1, value=f"College Management System – {title}")
    tc.font      = Font(name="Calibri", bold=True, size=13, color=XL_WHITE)
    tc.fill      = PatternFill("solid", fgColor=XL_NAVY)
    tc.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Sub-title row
    ws.merge_cells(start_row=2, start_column=1,
                   end_row=2,   end_column=len(headers))
    sc = ws.cell(row=2, column=1,
                 value=f"Generated: {datetime.datetime.now().strftime('%d %B %Y  %H:%M')}")
    sc.font      = Font(name="Calibri", italic=True, size=9, color="7f8c8d")
    sc.alignment = Alignment(horizontal="center")
    ws.row_dimensions[2].height = 16

    # Header row
    h_style = _xl_header_style()
    for ci, h in enumerate(headers, start=1):
        cell = ws.cell(row=3, column=ci, value=h)
        _xl_apply(cell, h_style)
    ws.row_dimensions[3].height = 20

    # Data rows
    for ri, row in enumerate(rows, start=4):
        even = (ri % 2 == 0)
        d_style = _xl_data_style(even)
        for ci, val in enumerate(row, start=1):
            cell = ws.cell(row=ri, column=ci,
                           value=val if val is not None else "")
            _xl_apply(cell, d_style)

    # Auto-fit column widths
    for ci, h in enumerate(headers, start=1):
        col_letter = get_column_letter(ci)
        max_len = max(
            len(str(h)),
            *[len(str(row[ci-1])) if row[ci-1] is not None else 0 for row in rows]
        ) if rows else len(str(h))
        ws.column_dimensions[col_letter].width = min(max_len + 4, 40)

    # Freeze header rows
    ws.freeze_panes = "A4"

    wb.save(path)


# ── CSV helper ────────────────────────────────────────────────────────────── #

def _build_csv(path, title, headers, rows):
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(["College Management System"])
        writer.writerow([title])
        writer.writerow([f"Generated: {datetime.datetime.now().strftime('%d %B %Y %H:%M')}"])
        writer.writerow([])          # blank separator
        writer.writerow(headers)
        for row in rows:
            writer.writerow([v if v is not None else "" for v in row])


# ════════════════════════════════════════════════════════════════════════════
#  PUBLIC REPORT GENERATORS
#  Each returns the saved file path.
#  fmt = 'pdf' | 'excel' | 'csv'
# ════════════════════════════════════════════════════════════════════════════

def generate_students_report(fmt="pdf") -> str:
    conn = get_connection()
    rows = conn.execute(
        """SELECT s.student_id, s.first_name||' '||s.last_name AS name,
                  s.age, s.gender, s.department, s.gpa,
                  s.city, u.email, s.phone, s.enrollment_year
           FROM students s
           JOIN users u ON s.user_id = u.user_id
           WHERE u.is_active = 1
           ORDER BY s.department, s.last_name"""
    ).fetchall()
    conn.close()

    headers = ["ID","Name","Age","Gender","Department","GPA",
               "City","Email","Phone","Enroll Year"]
    title   = "Student Records"
    path    = _save_path("students", fmt)

    if fmt == "pdf":
        _build_pdf(path, title, f"{len(rows)} active students", headers,
                   [tuple(r) for r in rows],
                   [1.2,3.5,0.9,1.1,2.5,1.1,2,3.8,2.4,1.3])
    elif fmt == "excel":
        _build_excel(path, "Students", title, headers, [tuple(r) for r in rows])
    else:
        _build_csv(path, title, headers, [tuple(r) for r in rows])
    return path


def generate_professors_report(fmt="pdf") -> str:
    conn = get_connection()
    rows = conn.execute(
        """SELECT p.prof_id, p.first_name||' '||p.last_name AS name,
                  p.age, p.gender, p.department, p.salary,
                  p.city, u.email, p.phone, p.joining_year
           FROM professors p
           JOIN users u ON p.user_id = u.user_id
           WHERE u.is_active = 1
           ORDER BY p.department, p.last_name"""
    ).fetchall()
    conn.close()

    headers = ["ID","Name","Age","Gender","Department","Salary",
               "City","Email","Phone","Join Year"]
    title   = "Professor Records"
    path    = _save_path("professors", fmt)

    if fmt == "pdf":
        _build_pdf(path, title, f"{len(rows)} active professors", headers,
                   [tuple(r) for r in rows],
                   [1.2,3.5,0.9,1.1,2.5,1.8,2,3.8,2.4,1.2])
    elif fmt == "excel":
        _build_excel(path, "Professors", title, headers, [tuple(r) for r in rows])
    else:
        _build_csv(path, title, headers, [tuple(r) for r in rows])
    return path


def generate_courses_report(fmt="pdf") -> str:
    conn = get_connection()
    rows = conn.execute(
        """SELECT c.course_id, c.course_name,
                  COALESCE(p.first_name||' '||p.last_name,'Unassigned') AS teacher,
                  c.credit_hours, c.semester, c.department,
                  c.fee, c.room_no, c.timing, c.status, c.year
           FROM courses c
           LEFT JOIN course_assignments ca ON ca.course_id = c.course_id
           LEFT JOIN professors p          ON p.prof_id    = ca.prof_id
           ORDER BY c.department, c.course_name"""
    ).fetchall()
    conn.close()

    headers = ["ID","Course","Teacher","Cr","Semester","Dept",
               "Fee","Room","Timing","Status","Year"]
    title   = "Course Records"
    path    = _save_path("courses", fmt)

    if fmt == "pdf":
        _build_pdf(path, title, f"{len(rows)} courses", headers,
                   [tuple(r) for r in rows],
                   [0.9,3.3,3.2,0.7,1.4,2.2,1.4,1.1,2,1.3,1.1])
    elif fmt == "excel":
        _build_excel(path, "Courses", title, headers, [tuple(r) for r in rows])
    else:
        _build_csv(path, title, headers, [tuple(r) for r in rows])
    return path


def generate_enrollments_report(fmt="pdf") -> str:
    conn = get_connection()
    rows = conn.execute(
        """SELECT s.student_id,
                  s.first_name||' '||s.last_name AS student,
                  s.department,
                  c.course_name,
                  c.credit_hours,
                  e.grade,
                  e.enrolled_at
           FROM enrollments e
           JOIN students s ON s.student_id = e.student_id
           JOIN courses  c ON c.course_id  = e.course_id
           ORDER BY s.department, s.last_name"""
    ).fetchall()
    conn.close()

    headers = ["Stu.ID","Student","Department","Course","Credits","Grade","Enrolled At"]
    title   = "Enrollment Records"
    path    = _save_path("enrollments", fmt)

    if fmt == "pdf":
        _build_pdf(path, title, f"{len(rows)} enrollments", headers,
                   [tuple(r) for r in rows],
                   [1.3,3.8,2.8,4,1.3,1.2,3])
    elif fmt == "excel":
        _build_excel(path, "Enrollments", title, headers, [tuple(r) for r in rows])
    else:
        _build_csv(path, title, headers, [tuple(r) for r in rows])
    return path


def generate_lab_report(fmt="pdf") -> str:
    conn = get_connection()
    rows = conn.execute(
        """SELECT la.lab_id,
                  la.first_name||' '||la.last_name AS attendant,
                  la.lab_name, la.shift,
                  la.city, u.email, la.phone, la.joining_year
           FROM lab_attendants la
           JOIN users u ON la.user_id = u.user_id
           WHERE u.is_active = 1
           ORDER BY la.lab_name"""
    ).fetchall()
    conn.close()

    headers = ["ID","Name","Lab","Shift","City","Email","Phone","Join Year"]
    title   = "Lab Attendant Records"
    path    = _save_path("lab_attendants", fmt)

    if fmt == "pdf":
        _build_pdf(path, title, f"{len(rows)} active attendants", headers,
                   [tuple(r) for r in rows],
                   [1.1,3.3,2.8,1.6,2,3.8,2.3,1.2])
    elif fmt == "excel":
        _build_excel(path, "Lab Attendants", title, headers, [tuple(r) for r in rows])
    else:
        _build_csv(path, title, headers, [tuple(r) for r in rows])
    return path


def generate_summary_report(fmt="pdf") -> str:
    conn    = get_connection()
    stats   = dict(conn.execute("SELECT * FROM vw_summary_stats").fetchone())
    dept_s  = conn.execute(
        "SELECT department, COUNT(*) FROM students GROUP BY department ORDER BY 2 DESC"
    ).fetchall()
    dept_p  = conn.execute(
        "SELECT department, COUNT(*) FROM professors GROUP BY department ORDER BY 2 DESC"
    ).fetchall()
    # Disabled users — always included in summary with "Disabled" label
    disabled = conn.execute(
        """SELECT u.user_id, u.email, u.role, u.created_at
           FROM users u
           WHERE u.is_active = 0 AND u.role != 'admin'
           ORDER BY u.role, u.email"""
    ).fetchall()
    conn.close()

    title = "System Summary Report"
    path  = _save_path("summary", fmt)

    kpi_headers = ["Metric", "Value"]
    kpi_rows    = [
        ("Active Students",       stats["active_students"]),
        ("Active Professors",     stats["active_professors"]),
        ("Active Lab Attendants", stats["active_lab_attendants"]),
        ("Disabled Accounts",     stats["disabled_users"]),
        ("Total Courses",         stats["total_courses"]),
        ("Total Enrollments",     stats["total_enrollments"]),
        ("Average Student GPA",   stats["avg_gpa"] or 0),
        ("Average Prof. Salary",  f"${stats['avg_salary'] or 0:,.2f}"),
    ]

    dis_headers = ["User ID", "Email", "Role", "Created At", "Status"]
    dis_rows    = [(r[0], r[1], r[2], r[3], "DISABLED") for r in disabled]

    if fmt == "pdf":
        doc   = SimpleDocTemplate(path, pagesize=A4,
                                  leftMargin=2*cm, rightMargin=2*cm,
                                  topMargin=1.5*cm, bottomMargin=1.5*cm)
        sty   = _pdf_styles()
        elems = _pdf_intro(sty, title, "All departments · All years")

        # KPI table
        kpi_data = [kpi_headers] + [[str(r[0]), str(r[1])] for r in kpi_rows]
        kpi_tbl  = Table(kpi_data, colWidths=[10*cm, 8*cm])
        kpi_tbl.setStyle(TableStyle([
            ("BACKGROUND",    (0,0),(-1,0),  PRIMARY),
            ("TEXTCOLOR",     (0,0),(-1,0),  WHITE),
            ("FONTNAME",      (0,0),(-1,0),  "Helvetica-Bold"),
            ("FONTSIZE",      (0,0),(-1,0),  10),
            ("ROWBACKGROUNDS",(0,1),(-1,-1), [WHITE, LIGHT_BG]),
            ("FONTSIZE",      (0,1),(-1,-1), 10),
            ("ALIGN",         (0,0),(-1,-1), "CENTER"),
            ("GRID",          (0,0),(-1,-1), 0.5, SECONDARY),
            ("TOPPADDING",    (0,0),(-1,-1), 8),
            ("BOTTOMPADDING", (0,0),(-1,-1), 8),
        ]))
        elems.append(kpi_tbl)
        elems.append(Spacer(1, 0.5*cm))

        # Dept breakdowns
        if dept_s:
            d_data = [["Department","Students"]] + [[r[0],str(r[1])] for r in dept_s]
            d_tbl  = Table(d_data, colWidths=[12*cm, 6*cm])
            d_tbl.setStyle(_pdf_table_style())
            elems += [Paragraph("Students by Department", sty["Normal"]),
                      Spacer(1,0.2*cm), d_tbl, Spacer(1,0.4*cm)]

        if dept_p:
            p_data = [["Department","Professors"]] + [[r[0],str(r[1])] for r in dept_p]
            p_tbl  = Table(p_data, colWidths=[12*cm, 6*cm])
            p_tbl.setStyle(_pdf_table_style())
            elems += [Paragraph("Professors by Department", sty["Normal"]),
                      Spacer(1,0.2*cm), p_tbl, Spacer(1,0.5*cm)]

        # Disabled users section
        if dis_rows:
            elems.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#c0392b")))
            elems.append(Spacer(1,0.2*cm))
            elems.append(Paragraph("Disabled Accounts", sty["Normal"]))
            elems.append(Spacer(1,0.2*cm))
            dis_data = [dis_headers] + [[str(v) for v in r] for r in dis_rows]
            dis_tbl  = Table(dis_data, colWidths=[1.5*cm, 5.5*cm, 3*cm, 4*cm, 2*cm])
            dis_style = _pdf_table_style()
            # Override: make STATUS column red text
            dis_style.add("TEXTCOLOR",   (4,1),(-1,-1), colors.HexColor("#c0392b"))
            dis_style.add("FONTNAME",    (4,1),(-1,-1), "Helvetica-Bold")
            dis_tbl.setStyle(dis_style)
            elems.append(dis_tbl)

        doc.build(elems)

    elif fmt == "excel":
        wb  = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Summary"

        # Title
        ws1.merge_cells("A1:B1")
        hc = ws1["A1"]
        hc.value     = f"College Management System – {title}"
        hc.font      = Font(bold=True, size=14, color=XL_WHITE)
        hc.fill      = PatternFill("solid", fgColor=XL_NAVY)
        hc.alignment = Alignment(horizontal="center")
        ws1.row_dimensions[1].height = 26

        h_style = _xl_header_style()
        for ci, h in enumerate(kpi_headers, 1):
            cell = ws1.cell(row=2, column=ci, value=h)
            _xl_apply(cell, h_style)

        for ri, (metric, value) in enumerate(kpi_rows, start=3):
            even = ri % 2 == 0
            for ci, val in enumerate([metric, value], 1):
                cell = ws1.cell(row=ri, column=ci, value=val)
                _xl_apply(cell, _xl_data_style(even))
        ws1.column_dimensions["A"].width = 28
        ws1.column_dimensions["B"].width = 20

        # Sheet 2: by department
        ws2 = wb.create_sheet("By Department")
        for ci, h in enumerate(["Department","Students","Professors"], 1):
            cell = ws2.cell(row=1, column=ci, value=h)
            _xl_apply(cell, h_style)
        dept_merged = {}
        for d, c in dept_s:
            dept_merged[d] = [d, c, 0]
        for d, c in dept_p:
            if d in dept_merged:
                dept_merged[d][2] = c
            else:
                dept_merged[d] = [d, 0, c]
        for ri, vals in enumerate(dept_merged.values(), start=2):
            for ci, v in enumerate(vals, 1):
                cell = ws2.cell(row=ri, column=ci, value=v)
                _xl_apply(cell, _xl_data_style(ri % 2 == 0))
        for col in ["A","B","C"]:
            ws2.column_dimensions[col].width = 22

        # Sheet 3: disabled users (always present, even if empty)
        ws3 = wb.create_sheet("Disabled Accounts")
        for ci, h in enumerate(dis_headers, 1):
            cell = ws3.cell(row=1, column=ci, value=h)
            # Red header for disabled sheet
            cell.font      = Font(name="Calibri", bold=True, color=XL_WHITE, size=10)
            cell.fill      = PatternFill("solid", fgColor="c0392b")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border    = _border

        if dis_rows:
            for ri, row in enumerate(dis_rows, start=2):
                for ci, val in enumerate(row, start=1):
                    cell = ws3.cell(row=ri, column=ci, value=val)
                    _xl_apply(cell, _xl_data_style(ri % 2 == 0))
                    # STATUS column: bold red text
                    if ci == 5:
                        cell.font = Font(name="Calibri", bold=True,
                                         color="c0392b", size=9)
        else:
            ws3.cell(row=2, column=1, value="No disabled accounts.")

        for col_letter, width in zip(["A","B","C","D","E"],
                                     [10, 35, 16, 22, 12]):
            ws3.column_dimensions[col_letter].width = width

        wb.save(path)

    else:  # csv
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            w.writerow(["College Management System"])
            w.writerow([title])
            w.writerow([f"Generated: {datetime.datetime.now().strftime('%d %B %Y %H:%M')}"])
            w.writerow([])
            w.writerow(kpi_headers)
            w.writerows(kpi_rows)
            w.writerow([])
            w.writerow(["--- Students by Department ---"])
            w.writerow(["Department", "Count"])
            w.writerows(dept_s)
            w.writerow([])
            w.writerow(["--- Professors by Department ---"])
            w.writerow(["Department", "Count"])
            w.writerows(dept_p)
            w.writerow([])
            w.writerow(["--- Disabled Accounts ---"])
            w.writerow(dis_headers)
            if dis_rows:
                w.writerows(dis_rows)
            else:
                w.writerow(["No disabled accounts."])

    return path